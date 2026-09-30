#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docs/ТЗ.md -> docs/ТЗ.docx

Нужен для того, чтобы техническое задание можно было править и пересылать
в Word, а не только читать в рабочем каталоге. Разбирает только те
конструкции, которые реально есть в ТЗ.md: заголовки, таблицы, списки,
абзацы, **жирный** и `код`.
"""

import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "ТЗ.md")
DST = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "ТЗ.docx")

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(par, text):
    """Текст с **жирным** и `кодом` -> набор runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            par.add_run(part)


def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_separator(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def main():
    with io.open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # --- таблица ---
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            header = split_row(stripped)
            body = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append(split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for n, text in enumerate(header):
                cell = table.rows[0].cells[n]
                cell.paragraphs[0].text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in body:
                cells = table.add_row().cells
                for n in range(len(header)):
                    text = row[n] if n < len(row) else ""
                    cells[n].paragraphs[0].text = ""
                    add_runs(cells[n].paragraphs[0], text)
            doc.add_paragraph()
            i = j
            continue

        # --- заголовки ---
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level, 4))
            i += 1
            continue

        # --- разделитель ---
        if stripped == "---":
            i += 1
            continue

        # --- нумерованный список ---
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Number")
            add_runs(par, m.group(2))
            i += 1
            continue

        # --- маркер списка, в том числе вложенный ---
        m = re.match(r"^(\s*)[*-]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            style_name = "List Bullet" if indent < 2 else "List Bullet 2"
            par = doc.add_paragraph(style=style_name)
            add_runs(par, m.group(2))
            i += 1
            continue

        # --- пустая строка ---
        if not stripped:
            i += 1
            continue

        # --- абзац: собираем следующие строки, пока не пустая ---
        chunk = [stripped]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (not nxt or nxt.startswith("#") or nxt.startswith("|")
                    or nxt == "---" or re.match(r"^(\s*)[*-]\s+", lines[j])
                    or re.match(r"^\d+\.\s+", nxt)):
                break
            chunk.append(nxt)
            j += 1
        par = doc.add_paragraph()
        add_runs(par, " ".join(chunk))
        i = j

    doc.save(DST)
    print("сохранено:", DST)


if __name__ == "__main__":
    sys.exit(main())
